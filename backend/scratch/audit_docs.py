import hashlib
import os
import pypdf
import pptx
import docx

dir_path = "data/phd-admission"
if not os.path.exists(dir_path):
    dir_path = "../data/phd-admission"

files = os.listdir(dir_path)
print("Files found in", os.path.abspath(dir_path), ":", files)

for f in sorted(files):
    p = os.path.join(dir_path, f)
    size = os.path.getsize(p)
    with open(p, "rb") as fp:
        sha = hashlib.sha256(fp.read()).hexdigest()
    print(f"\n========================================================")
    print(f"FILE: {f}")
    print(f"Size: {size:,} bytes")
    print(f"SHA256: {sha}")

    if f.endswith(".pdf"):
        reader = pypdf.PdfReader(p)
        total_pages = len(reader.pages)
        print(f"Format: PDF | Total Pages: {total_pages}")
        empty_pages = []
        page_lengths = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            txt_stripped = txt.strip()
            page_lengths.append(len(txt_stripped))
            if len(txt_stripped) == 0:
                empty_pages.append(i + 1)
        print(f"Empty pages: {empty_pages}")
        print(f"Average characters per page: {sum(page_lengths)/total_pages:.1f}")
        print("Page 1 snippet:\n", (reader.pages[0].extract_text() or "")[:400].strip())
        print("Page 2 snippet:\n", (reader.pages[1].extract_text() or "")[:400].strip())

    elif f.endswith(".pptx"):
        prs = pptx.Presentation(p)
        total_slides = len(prs.slides)
        print(f"Format: PPTX | Total Slides: {total_slides}")
        empty_slides = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text.strip())
            full = " ".join(t for t in texts if t)
            if not full:
                empty_slides.append(i + 1)
            else:
                title = slide.shapes.title.text if slide.shapes.title and slide.shapes.title.has_text_frame else (texts[0] if texts else "No title")
                if i < 5 or i >= total_slides - 2:
                    print(f"  Slide {i+1}: {title[:80]} ({len(full)} chars)")
        print(f"Empty slides: {empty_slides}")

    elif f.endswith(".docx"):
        doc = docx.Document(p)
        print(f"Format: DOCX | Total Paragraphs: {len(doc.paragraphs)} | Total Tables: {len(doc.tables)}")
        non_empty = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        print(f"Non-empty paragraphs: {len(non_empty)}")
        for i, text in enumerate(non_empty[:6]):
            print(f"  Para {i+1}: {text[:100]}")
